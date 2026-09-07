from fastapi import FastAPI, Request ,File,UploadFile
from fastapi.responses import RedirectResponse,HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from db import SessionLocal 
from ai import analyse_resuma
from typing import Optional
from starlette.middleware.sessions import SessionMiddleware
from io import BytesIO
import pdfplumber
import json
import PyPDF2
import docx
import models


app = FastAPI()

app.add_middleware(SessionMiddleware, secret_key="super-secret-key")
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

# app = FastAPI()
# templates = Jinja2Templates(directory="templates")

@app.api_route("/", methods=["GET", "POST"])
async def signup(request: Request):
    if request.method == "POST":
        db = SessionLocal()
        form = await request.form()
        email = form.get("Email")
        password = form.get("Password")
        existing_user = db.query(models.User).filter_by(email_id=email).first()

        if existing_user:
            return templates.TemplateResponse(request,
                "signup.html",
                {"request": request, "error": "User already exists"},status_code=200
            )

        user = models.User(email_id=email, password=password)
        db.add(user)
        db.commit()

        return RedirectResponse("/login", status_code=302)

    return templates.TemplateResponse(request,"signup.html", {"request": request})

# login
@app.api_route("/login",methods=["POST","GET"],response_class=HTMLResponse)
async def login(request:Request):
    db = SessionLocal()
    if request.method == "POST":
        form = await request.form()
        email = form.get("Email")
        password = form.get("Password")
        user = db.query(models.User).filter_by(email_id=email).first()
        if user :
            request.session["user"] =user.email_id
            return RedirectResponse("/dashboard",status_code=302)
        else:
            return "Invalide User"

    return templates.TemplateResponse(request,"login.html",{"request":request})

# Dashboard
@app.api_route("/dashboard",methods=["POST","GET"],response_class=HTMLResponse)
async def dashboard(request:Request,file: Optional[UploadFile] = None):
    db = SessionLocal()
    if "user" not in request.session:
        return RedirectResponse("/login")
    result = None
    if request.method == "POST":
        form = await request.form()
        # contents = await file.read()
        user_goal = form.get("role")
        user_resuma_text = form.get("resuma")
        # file : UploadFile = File(None)
        # if file and file.filename :   # ✅ check UploadFile, not contents
        #     contents = await file.read()

        # # if file and file.filename != "":
        #     if contents and contents.filename != "":
        #         if file.filename.endswith (".pdf"):
        #             try:
        #                 # pdf_reder = PyPDF2.PdfReader(file.file)
        #                 pdf_reder = PyPDF2.PdfReader(BytesIO(contents))
        #                 text = ""
        #                 for page in pdf_reder.pages:
        #                     text+=page.extract_text() or ""
        #                 user_resuma_text = text
        #             except Exception as e:
        #                 result = {"error": f"Error reading PDF file:{e}"}
        #         # elif file.filename.endswith(".docx"):
        #         elif contents.filename.endswith(".docx"):
        #             try:
        #                 # doc = docx.Document(file.file)
        #                 doc = docx.Document(BytesIO(contents))
        #                 text = ""
        #                 for parge in doc.paragraphs:
        #                     text += parge.text or ""
        #                 user_resuma_text = text
        #             except Exception as e:
        #                 result = {"error": f"Error reading DOCX file: {e}"}
        if file and file.filename:
            contents = await file.read()

            if not contents:
                result = {"error": "Uploaded file is empty"}
            elif file.filename.endswith(".pdf"):
                try:
                    text = ""
                    with pdfplumber.open(BytesIO(contents)) as pdf:
                        for page in pdf.pages:
                            text += page.extract_text() or ""
                    user_resuma_text = text

                    #### use diffrent things ####
            #         pdf_reader = PyPDF2.PdfReader(BytesIO(contents))
            #         text = "".join(page.extract_text() or "" for page in pdf_reader.pages)
            #         user_resuma_text = text

                except Exception as e:
                    result = {"error": f"Error reading PDF file: {e}"}


            elif file.filename.endswith(".docx"):
                try:
                    doc = docx.Document(BytesIO(contents))
                    text = "".join(p.text or "" for p in doc.paragraphs)
                    user_resuma_text = text
                except Exception as e:
                    result = {"error": f"Error reading DOCX file: {e}"}


        if user_goal and user_resuma_text:
            # Process the user's goal and resume text here
            try:
                clean_text = " ".join(user_resuma_text.split())
                combined_text = f"Role: {user_goal}\nResume:\n{clean_text}"
                # result =analyse_resuma(user_goal, user_resuma_text)
                result =analyse_resuma(user_goal, combined_text)
                # save to db 
                db = SessionLocal()
                user = db.query(models.User).filter_by(email_id=request.session["user"]).first()
                report = models.Report(
                    u_id = user.id,
                    resuma_text = user_resuma_text,
                    result= json.dumps(result)
                    # goal = user_goal,
                    # skills = json.dumps(result.get("skills", [])),
                    # missing_skills = json.dumps(result.get("missing_skills", [])),
                    # roadmap = json.dumps(result.get("roadmap", [])),
                    # interview_questions = json.dumps(result.get("interview_questions", []))
                )
                db.add(report)
                db.commit()
                return templates.TemplateResponse(request,
                                                  "dashboard.html",{"request":request,
                                                    "user":request.session["user"],
                                                    "result":result
                                                    })            
            except Exception as e:
                result = {"error": f"AI Error processing data: {str(e)}"}
    return templates.TemplateResponse(request,"dashboard.html",{
        "request":request,
        "user":request.session["user"],
        "result":result
    })

# history
@app.api_route("/history",methods=["POST","GET"],response_class=HTMLResponse)
async def history(request:Request):
    session = request.session
    if "user" not in session:
        return RedirectResponse("/login")
    db = SessionLocal()
    user = db. query(models.User).filter_by(email_id=session["user"]).first()
    reports = db.query(models.Report).filter_by(u_id=user.id).all()

    ### conver the json String >dict
    parshed_reports =[]
    for r in reports:
            try:
                parshed_result = json.loads(r.result)
            except Exception as e:
                parshed_result=[]
            parshed_reports.append({
                "resuma":r.resuma_text ,
                "result":parshed_result
            })
    return templates.TemplateResponse(request,"history.html",{"request":request,"reports":parshed_reports})

@app.api_route("/logout",response_class=HTMLResponse)
async def logout(request:Request):
    session = request.session
    session.pop("user",None)
    return RedirectResponse("/login")